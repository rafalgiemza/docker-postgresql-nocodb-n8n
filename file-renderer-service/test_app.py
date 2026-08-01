"""Integration test for the /render and /render-docx endpoints - no NocoDB, no network.
app.py has zero NocoDB knowledge now, so this just drives the real FastAPI
handlers with in-memory template + data payloads.

Run:  pytest file-renderer-service/test_app.py -v
"""
import asyncio
import io
import json

from pptx import Presentation
from pptx.util import Inches
from docx import Document
from starlette.datastructures import UploadFile

import app


def _template_pptx_bytes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.paragraphs[0].add_run().text = "{{lead.contact_name}}"

    repeat = prs.slides.add_slide(prs.slide_layouts[6])
    box2 = repeat.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box2.text_frame.paragraphs[0].add_run().text = "{{participant.full_name}}"
    repeat.notes_slide.notes_text_frame.text = "repeat:participant"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _call_render(template_bytes, data):
    upload = UploadFile(file=io.BytesIO(template_bytes), filename="template.pptx")
    response = asyncio.run(app.render(template=upload, data=json.dumps(data)))
    return response


def test_render_returns_rendered_pptx_and_no_warnings():
    data = {
        "lead": {"contact_name": "Ala"},
        "participant": [{"full_name": "Basia"}, {"full_name": "Czesiek"}],
    }
    response = _call_render(_template_pptx_bytes(), data)

    assert response.media_type == app.PPTX_MIME
    warnings = json.loads(response.headers["X-Warnings"])
    assert warnings == []

    rendered = Presentation(io.BytesIO(response.body))
    slides = list(rendered.slides)
    assert len(slides) == 3  # 1 untouched + 2 repeats
    assert slides[0].shapes[0].text_frame.text == "Ala"
    assert [s.shapes[0].text_frame.text for s in slides[1:]] == ["Basia", "Czesiek"]


def test_render_surfaces_warnings_in_header():
    data = {"lead": {}, "participant": []}
    response = _call_render(_template_pptx_bytes(), data)

    warnings = json.loads(response.headers["X-Warnings"])
    assert any("lead.contact_name" in w for w in warnings)
    assert any("repeat:participant" in w and "dropped" in w for w in warnings)


def _template_docx_bytes():
    """Create a simple DOCX template with placeholders."""
    doc = Document()
    doc.add_paragraph("Contact: {{lead.contact_name}}")
    doc.add_paragraph("Email: {{lead.email}}")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].paragraphs[0].text = "Name"
    table.rows[0].cells[1].paragraphs[0].text = "{{person.name}}"
    table.rows[1].cells[0].paragraphs[0].text = "Phone"
    table.rows[1].cells[1].paragraphs[0].text = "{{person.phone}}"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _call_render_docx(template_bytes, data):
    upload = UploadFile(file=io.BytesIO(template_bytes), filename="template.docx")
    response = asyncio.run(app.render_document(template=upload, data=json.dumps(data)))
    return response


def test_render_docx_returns_rendered_docx_and_no_warnings():
    data = {
        "lead": {"contact_name": "Ala", "email": "ala@example.com"},
        "person": {"name": "Basia", "phone": "123456789"},
    }
    response = _call_render_docx(_template_docx_bytes(), data)

    assert response.media_type == app.DOCX_MIME
    warnings = json.loads(response.headers["X-Warnings"])
    assert warnings == []

    rendered = Document(io.BytesIO(response.body))
    assert rendered.paragraphs[0].text == "Contact: Ala"
    assert rendered.paragraphs[1].text == "Email: ala@example.com"

    table = rendered.tables[0]
    assert table.rows[0].cells[1].text == "Basia"
    assert table.rows[1].cells[1].text == "123456789"


def test_render_docx_surfaces_warnings_in_header():
    data = {"lead": {}, "person": {}}
    response = _call_render_docx(_template_docx_bytes(), data)

    warnings = json.loads(response.headers["X-Warnings"])
    assert any("lead.contact_name" in w for w in warnings)
    assert any("lead.email" in w for w in warnings)
