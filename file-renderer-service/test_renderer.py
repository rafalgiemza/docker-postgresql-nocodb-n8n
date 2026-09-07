"""Offline unit tests for renderer.py — no NocoDB, no network.

Run:  pytest file-renderer-service/test_renderer.py -v
"""
import io

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from renderer import render_pptx


def _new_prs():
    return Presentation()


def _add_slide(prs, runs_text, notes=None):
    """Blank-layout slide with one textbox; runs_text is a list of run strings
    (len > 1 simulates PowerPoint splitting a placeholder across runs)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    para = box.text_frame.paragraphs[0]
    for text in runs_text:
        run = para.add_run()
        run.text = text
    if notes is not None:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _render(prs, data, warnings):
    return Presentation(io.BytesIO(render_pptx(_bytes(prs), data, warnings)))


def _bytes(prs):
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _textbox_text(slide):
    return slide.shapes[0].text_frame.text


def _para_children(slide, para_idx=0):
    """Ordered (kind, text, bold) tuples for a paragraph's a:r/a:br children -
    text_frame.text alone can't tell a line break from nothing, and .runs
    skips a:br entirely."""
    para = slide.shapes[0].text_frame.paragraphs[para_idx]
    out = []
    for el in para._p:
        if el.tag == qn("a:r"):
            t = el.find(qn("a:t"))
            rpr = el.find(qn("a:rPr"))
            bold = rpr.get("b") if rpr is not None else None
            out.append(("r", t.text if t is not None else "", bold))
        elif el.tag == qn("a:br"):
            out.append(("br", None, None))
    return out


def test_simple_placeholder_resolves():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    warnings = []
    out = _render(prs, {"lead": {"contact_name": "Ala"}}, warnings)
    assert _textbox_text(out.slides[0]) == "Ala"
    assert warnings == []


def test_missing_value_becomes_empty_string_and_warns():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    warnings = []
    out = _render(prs, {"lead": {}}, warnings)
    assert _textbox_text(out.slides[0]) == ""
    assert any("lead.contact_name" in w for w in warnings)


def test_placeholder_split_across_runs_collapses_to_first_run():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.", "contact_name}}"])
    warnings = []
    out = _render(prs, {"lead": {"contact_name": "Ala"}}, warnings)
    para = out.slides[0].shapes[0].text_frame.paragraphs[0]
    assert [r.text for r in para.runs] == ["Ala", ""]


def test_repeat_marker_name_is_generic_not_hardcoded():
    """repeat:<anything> works - the name is just a data key + placeholder
    prefix, nothing about "participant"/"testimonial" is special-cased."""
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    _add_slide(prs, ["{{goal.title}}"], notes="repeat:goal")
    warnings = []
    data = {
        "lead": {"contact_name": "Ala"},
        "goal": [{"title": "A"}, {"title": "B"}, {"title": "C"}],
    }
    out = _render(prs, data, warnings)
    slides = list(out.slides)
    assert len(slides) == 4  # 1 untouched + 3 repeats
    assert _textbox_text(slides[0]) == "Ala"
    assert [_textbox_text(s) for s in slides[1:]] == ["A", "B", "C"]
    assert warnings == []


def test_repeat_with_empty_list_drops_slide_and_warns():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    _add_slide(prs, ["{{participant.full_name}}"], notes="repeat:participant")
    warnings = []
    data = {"lead": {"contact_name": "Ala"}, "participant": []}
    out = _render(prs, data, warnings)
    assert len(out.slides) == 1
    assert _textbox_text(out.slides[0]) == "Ala"
    assert any("repeat:participant" in w and "dropped" in w for w in warnings)


def test_repeat_missing_key_in_data_drops_slide_and_warns():
    prs = _new_prs()
    _add_slide(prs, ["{{testimonial.client_name}}"], notes="repeat:testimonial")
    warnings = []
    out = _render(prs, {}, warnings)  # no "testimonial" key at all
    assert len(out.slides) == 0
    assert any("repeat:testimonial" in w and "dropped" in w for w in warnings)


def test_repeat_testimonial_uses_matching_prefix():
    prs = _new_prs()
    _add_slide(prs, ["{{testimonial.client_name}}"], notes="repeat:testimonial")
    warnings = []
    data = {"testimonial": [{"client_name": "Firma X"}]}
    out = _render(prs, data, warnings)
    assert _textbox_text(list(out.slides)[0]) == "Firma X"
    assert warnings == []


def test_nested_value_in_repeat_item_resolves():
    prs = _new_prs()
    _add_slide(prs, ["{{participant.full_name}} {{participant.a.o}}"],
               notes="repeat:participant")
    warnings = []
    data = {"participant": [{"full_name": "Basia", "a": {"o": "B2"}}]}
    out = _render(prs, data, warnings)
    assert _textbox_text(list(out.slides)[0]) == "Basia B2"
    assert warnings == []


def test_participant_without_nested_value_warns():
    prs = _new_prs()
    _add_slide(prs, ["{{participant.a.o}}"], notes="repeat:participant")
    warnings = []
    data = {"participant": [{"full_name": "Basia"}]}  # no "a" key
    _render(prs, data, warnings)
    assert any("participant.a.o" in w for w in warnings)


def test_repeat_item_keys_are_also_available_without_prefix():
    """Inside a repeat slide the item's own keys are lifted to the top level,
    so a template may write {{a.o}} instead of {{participant.a.o}}."""
    prs = _new_prs()
    _add_slide(prs, ["{{participant.full_name}} / {{full_name}} / {{a.o}} / {{assessment.position}}"],
               notes="repeat:participant")
    warnings = []
    data = {"participant": [{
        "full_name": "Basia",
        "a": {"o": "B2"},
        "assessment": {"position": "HR Manager"},
    }]}
    out = _render(prs, data, warnings)
    assert _textbox_text(list(out.slides)[0]) == "Basia / Basia / B2 / HR Manager"
    assert warnings == []


def test_lifted_item_keys_shadow_top_level_data_on_that_slide():
    prs = _new_prs()
    _add_slide(prs, ["{{title}}|{{lead.title}}"], notes="repeat:row")
    warnings = []
    data = {"lead": {"title": "z leada"}, "title": "globalny",
            "row": [{"title": "z elementu"}]}
    out = _render(prs, data, warnings)
    assert _textbox_text(list(out.slides)[0]) == "z elementu|z leada"


def test_bold_span_becomes_a_real_bold_run_not_literal_asterisks():
    prs = _new_prs()
    _add_slide(prs, ["{{module.goal_statement}}"])
    warnings = []
    data = {"module": {"goal_statement": "Zwykly tekst **pogrubiony fragment** koniec."}}
    out = _render(prs, data, warnings)
    assert _para_children(out.slides[0]) == [
        ("r", "Zwykly tekst ", "0"),
        ("r", "pogrubiony fragment", "1"),
        ("r", " koniec.", "0"),
    ]
    assert warnings == []


def test_br_and_newline_become_line_break_not_literal_text():
    prs = _new_prs()
    _add_slide(prs, ["{{module.goal_statement}}"])
    warnings = []
    data = {"module": {"goal_statement": "Pierwszy.<br>Drugi.\nTrzeci."}}
    out = _render(prs, data, warnings)
    assert _para_children(out.slides[0]) == [
        ("r", "Pierwszy.", "0"),
        ("br", None, None),
        ("r", "Drugi.", "0"),
        ("br", None, None),
        ("r", "Trzeci.", "0"),
    ]
    # python-pptx represents each a:br as "\x0b" (soft line break) when
    # joining paragraph text - a real break marker, not literal "<br>"/"\n".
    assert _textbox_text(out.slides[0]) == "Pierwszy.\x0bDrugi.\x0bTrzeci."


def test_plain_value_without_markup_stays_a_single_run():
    """No **/<br>/\\n in the resolved value -> untouched fast path, same as
    before this feature (no gratuitous run-splitting for ordinary text)."""
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    warnings = []
    out = _render(prs, {"lead": {"contact_name": "Ala"}}, warnings)
    para = out.slides[0].shapes[0].text_frame.paragraphs[0]
    assert len(para.runs) == 1
    assert para.runs[0].text == "Ala"


def test_rich_text_preserves_surrounding_plain_runs_and_their_order():
    prs = _new_prs()
    _add_slide(prs, ["Cel: ", "{{module.goal_statement}}", " (koniec)"])
    warnings = []
    data = {"module": {"goal_statement": "**A**\nB"}}
    out = _render(prs, data, warnings)
    assert _textbox_text(out.slides[0]) == "Cel: A\x0bB (koniec)"
    kinds = [k for k, _, _ in _para_children(out.slides[0])]
    assert kinds == ["r", "r", "br", "r", "r"]


def test_rich_text_when_placeholder_split_across_runs():
    """Slow path (placeholder text itself split across runs by PowerPoint) -
    the collapsed-into-first-run value still gets bold/break treatment."""
    prs = _new_prs()
    _add_slide(prs, ["{{module.goal_", "statement}}"])
    warnings = []
    data = {"module": {"goal_statement": "**Cel**\nszczegoly"}}
    out = _render(prs, data, warnings)
    assert _textbox_text(out.slides[0]) == "Cel\x0bszczegoly"
    kinds = [k for k, _, _ in _para_children(out.slides[0])]
    # trailing "r" is the original second run, blanked (not removed) - same
    # pre-existing collapse behavior as test_placeholder_split_across_runs_*.
    assert kinds == ["r", "br", "r", "r"]
