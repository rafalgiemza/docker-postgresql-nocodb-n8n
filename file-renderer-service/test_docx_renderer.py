"""Offline unit tests for docx_renderer.py — no NocoDB, no network.

Run:  pytest file-renderer-service/test_docx_renderer.py -v
"""
import io

from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.text.font import Font

from docx_renderer import render_docx


def _new_doc():
    return Document()


def _add_paragraph(doc, text):
    """Add a paragraph with a single run."""
    p = doc.add_paragraph(text)
    return p


def _add_table_with_text(doc, rows, cols, texts):
    """Add a table with text in cells. texts is a flat list."""
    table = doc.add_table(rows=rows, cols=cols)
    i = 0
    for row in table.rows:
        for cell in row.cells:
            if i < len(texts):
                cell.paragraphs[0].text = texts[i]
                i += 1
    return table


def _render(doc, data, warnings):
    """Render doc and return Document object."""
    buf = io.BytesIO()
    doc.save(buf)
    rendered_bytes = render_docx(buf.getvalue(), data, warnings)
    return Document(io.BytesIO(rendered_bytes))


def test_simple_placeholder_in_paragraph():
    doc = _new_doc()
    _add_paragraph(doc, "Hello {{lead.name}}")
    warnings = []
    out = _render(doc, {"lead": {"name": "Alice"}}, warnings)
    assert out.paragraphs[0].text == "Hello Alice"
    assert warnings == []


def test_missing_placeholder_becomes_empty_and_warns():
    doc = _new_doc()
    _add_paragraph(doc, "Hello {{lead.name}}")
    warnings = []
    out = _render(doc, {"lead": {}}, warnings)
    assert out.paragraphs[0].text == "Hello "
    assert any("lead.name" in w for w in warnings)


def test_nested_placeholder():
    doc = _new_doc()
    _add_paragraph(doc, "Contact: {{lead.person.email}}")
    warnings = []
    out = _render(doc, {"lead": {"person": {"email": "alice@example.com"}}}, warnings)
    assert out.paragraphs[0].text == "Contact: alice@example.com"
    assert warnings == []


def test_placeholder_in_table_cell():
    doc = _new_doc()
    _add_table_with_text(doc, 2, 2, ["Name", "{{person.name}}", "Email", "{{person.email}}"])
    warnings = []
    out = _render(doc, {"person": {"name": "Bob", "email": "bob@ex.com"}}, warnings)

    table = out.tables[0]
    cells = [cell.text for row in table.rows for cell in row.cells]
    assert cells == ["Name", "Bob", "Email", "bob@ex.com"]
    assert warnings == []


def test_multiple_placeholders_in_one_paragraph():
    doc = _new_doc()
    _add_paragraph(doc, "{{title}}: {{value}} ({{unit}})")
    warnings = []
    out = _render(doc, {"title": "Price", "value": "100", "unit": "PLN"}, warnings)
    assert out.paragraphs[0].text == "Price: 100 (PLN)"
    assert warnings == []


def test_no_placeholder_leaves_text_unchanged():
    doc = _new_doc()
    _add_paragraph(doc, "This is plain text")
    warnings = []
    out = _render(doc, {}, warnings)
    assert out.paragraphs[0].text == "This is plain text"
    assert warnings == []


def test_float_converted_to_int_if_whole():
    doc = _new_doc()
    _add_paragraph(doc, "Value: {{amount}}")
    warnings = []
    out = _render(doc, {"amount": 42.0}, warnings)
    assert out.paragraphs[0].text == "Value: 42"
    assert warnings == []


def test_float_with_decimal_stays_float():
    doc = _new_doc()
    _add_paragraph(doc, "Value: {{amount}}")
    warnings = []
    out = _render(doc, {"amount": 42.5}, warnings)
    assert out.paragraphs[0].text == "Value: 42.5"
    assert warnings == []


def _para_runs(paragraph):
    """(text, Font) pairs for every w:r in the paragraph, INCLUDING ones
    nested inside a w:hyperlink (paragraph.runs skips those)."""
    out = []
    for el in paragraph._p.iter():
        if el.tag == qn("w:r"):
            t = el.find(qn("w:t"))
            out.append((t.text if t is not None else "", Font(el)))
    return out


def test_bold_span_becomes_a_real_bold_run_not_literal_asterisks():
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "przed **pogrubione** po"}, warnings)
    runs = _para_runs(out.paragraphs[0])
    assert [(t, f.bold) for t, f in runs] == [
        ("przed ", False), ("pogrubione", True), (" po", False),
    ]
    assert warnings == []


def test_italic_span_becomes_a_real_italic_run():
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "przed _kursywa_ po"}, warnings)
    runs = _para_runs(out.paragraphs[0])
    assert [(t, f.italic) for t, f in runs] == [
        ("przed ", False), ("kursywa", True), (" po", False),
    ]
    assert warnings == []


def test_underline_span_becomes_a_real_underline_run():
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "przed <u>podkreslenie</u> po"}, warnings)
    runs = _para_runs(out.paragraphs[0])
    assert [(t, f.underline) for t, f in runs] == [
        ("przed ", False), ("podkreslenie", True), (" po", False),
    ]
    assert warnings == []


def test_single_word_strikethrough_becomes_a_real_strike_run():
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "przed ~slowo~ po"}, warnings)
    runs = _para_runs(out.paragraphs[0])
    assert [(t, f.strike) for t, f in runs] == [
        ("przed ", False), ("slowo", True), (" po", False),
    ]
    assert warnings == []


def test_multi_word_strikethrough_stays_literal_text():
    """Same NocoDB-syntax rule as PPTX: ~word~ (single tilde, single word),
    NOT GFM's ~~word~~ - a multi-word span is left untouched so an
    unrelated stray "~" in prose (e.g. "~10 godzin") can't misfire."""
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "~dwa slowa~ bez zmian"}, warnings)
    assert out.paragraphs[0].text == "~dwa slowa~ bez zmian"
    assert warnings == []


def test_inline_code_gets_monospace_font():
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "wpisz `kod` tutaj"}, warnings)
    runs = _para_runs(out.paragraphs[0])
    assert [(t, f.name) for t, f in runs] == [
        ("wpisz ", None), ("kod", "Consolas"), (" tutaj", None),
    ]
    assert warnings == []


def test_link_becomes_a_real_hyperlink_run():
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "zobacz [tutaj](https://example.com/oferta)"}, warnings)
    para_el = out.paragraphs[0]._p
    hyperlinks = [el for el in para_el.iter() if el.tag == qn("w:hyperlink")]
    assert len(hyperlinks) == 1
    rid = hyperlinks[0].get(qn("r:id"))
    assert out.part.rels[rid].target_ref == "https://example.com/oferta"
    run_text = hyperlinks[0].find(qn("w:r")).find(qn("w:t")).text
    assert run_text == "tutaj"
    assert warnings == []


def test_bold_italic_underline_combine_on_one_run():
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "**_<u>wszystko naraz</u>_**"}, warnings)
    runs = _para_runs(out.paragraphs[0])
    assert len(runs) == 1
    text, f = runs[0]
    assert (text, f.bold, f.italic, f.underline) == ("wszystko naraz", True, True, True)
    assert warnings == []


def test_br_and_newline_become_line_break_not_literal_text():
    """DOCX had ZERO rich-text handling before Phase 1 - <br>/\\n now become
    a real w:br (soft line break within the same paragraph), not literal
    text, matching what PPTX already did."""
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "Pierwszy.<br>Drugi.\nTrzeci."}, warnings)
    assert out.paragraphs[0].text == "Pierwszy.\nDrugi.\nTrzeci."
    para_el = out.paragraphs[0]._p
    br_count = sum(1 for el in para_el.iter() if el.tag == qn("w:br"))
    assert br_count == 2
    assert warnings == []


def test_plain_value_with_trailing_space_stays_a_single_run():
    """Regression guard: mistune strips block-edge whitespace internally -
    richtext.segments() must re-attach it so an everyday value like "Ala "
    doesn't gratuitously split into two runs (or lose the space)."""
    doc = _new_doc()
    _add_paragraph(doc, "{{v}}")
    warnings = []
    out = _render(doc, {"v": "Ala "}, warnings)
    assert len(out.paragraphs[0].runs) == 1
    assert out.paragraphs[0].runs[0].text == "Ala "
    assert warnings == []
