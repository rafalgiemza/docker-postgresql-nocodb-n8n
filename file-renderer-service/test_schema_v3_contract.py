"""Pins down that the renderer handles the data shape produced after the
schema v3 migration (docs/archive/fable/nocodb_crm_schema_v3.md) WITHOUT any code change -
and encodes the two limitations that shape imposes on how n8n must assemble
`data`. No NocoDB, no network.

Run:  pytest file-renderer-service/test_schema_v3_contract.py -v
"""
import io

from pptx import Presentation
from pptx.util import Inches

from renderer import render_pptx


def _add_slide(prs, text, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.paragraphs[0].add_run().text = text
    if notes is not None:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _render(prs, data, warnings):
    buf = io.BytesIO()
    prs.save(buf)
    return Presentation(io.BytesIO(render_pptx(buf.getvalue(), data, warnings)))


def _text(slide):
    return slide.shapes[0].text_frame.text


# The shape n8n (W9) will send once schema v3 is in place. `module` is
# deliberately FLAT - one row per (participant, module) pair - see
# test_nested_repeat_is_not_supported below for why.
DATA = {
    "lead": {"contact_name": "Piotr Zieliński", "type": "B2B"},
    "company": {"name": "ACME", "communication_processes": "demo dla klientów US"},
    "meeting": {"goals": "pewność w rozmowach", "challenges": "client demos",
                "business_context": "software house",
                "communication_situations": "stand-upy"},
    "offer": {"date": "01.08.2026", "price": 20400, "hours": 60,
              "variant": "standard"},
    "participant": [
        {"full_name": "Piotr", "position": "Dev",
         "a": {"o": "B2.4"},
         "assessment": {"needs_summary": "demo", "strengths": "słownictwo",
                        "gaps": "articles"},
         "recommendation": {"type": "business_english", "headline": "BE + Skills",
                            "rationale": "bo demo", "priority": "wysoki"}},
        {"full_name": "Anna", "position": "PO",
         "a": {"o": "C1.0"},
         "assessment": {"needs_summary": "negocjacje", "strengths": "płynność",
                        "gaps": "rejestr"},
         "recommendation": {"type": "skills_only", "headline": "Negocjacje",
                            "rationale": "bo kontrakty", "priority": "średni"}},
    ],
    "module": [
        {"sort_order": 1, "hours": 40, "mode": "1-1", "title": "Business English",
         "goal_statement": "swobodne demo", "participant_name": "Piotr"},
        {"sort_order": 2, "hours": 20, "mode": "w_parach", "title": "Negocjacje",
         "goal_statement": "kontrakty", "participant_name": "Anna"},
    ],
    "testimonial": [{"title": "Case", "client_name": "X", "position": "CTO",
                     "content": "ok"}],
}


def test_module_repeat_needs_no_code_change():
    """`recommendation_items` + `training_modules` -> a repeat:module slide."""
    prs = Presentation()
    _add_slide(prs, "ETAP {{module.sort_order}}: {{module.title}} "
                    "{{module.hours}}h ({{module.mode}}) dla {{module.participant_name}}",
               notes="repeat:module")
    warnings = []
    out = _render(prs, DATA, warnings)
    assert [_text(s) for s in out.slides] == [
        "ETAP 1: Business English 40h (1-1) dla Piotr",
        "ETAP 2: Negocjacje 20h (w_parach) dla Anna",
    ]
    assert warnings == []


def test_nested_recommendation_and_assessment_on_participant_slide():
    """v3 nests recommendation/assessment under each participant; both the
    prefixed and the lifted form must resolve."""
    prs = Presentation()
    _add_slide(prs, "{{participant.full_name}}: {{participant.recommendation.headline}}"
                    " / {{recommendation.type}} / {{assessment.gaps}} / {{a.o}}",
               notes="repeat:participant")
    warnings = []
    out = _render(prs, DATA, warnings)
    assert [_text(s) for s in out.slides] == [
        "Piotr: BE + Skills / business_english / articles / B2.4",
        "Anna: Negocjacje / skills_only / rejestr / C1.0",
    ]
    assert warnings == []


def test_meeting_and_company_fields_resolve_globally():
    """The split-out meetings.* fields (v3 §2) and companies.* (v3 §10)."""
    prs = Presentation()
    _add_slide(prs, "{{meeting.goals}} | {{meeting.challenges}} | "
                    "{{company.communication_processes}} | {{offer.hours}}h")
    warnings = []
    out = _render(prs, DATA, warnings)
    assert _text(list(out.slides)[0]) == (
        "pewność w rozmowach | client demos | demo dla klientów US | 60h")
    assert warnings == []


def test_participant_prefix_on_a_module_slide_is_a_trap():
    """On a repeat:module slide `{{participant.x}}` points at the GLOBAL
    participant LIST, not at a person - hence the flattened
    `module.participant_name`. Warns instead of silently rendering nothing
    useful."""
    prs = Presentation()
    _add_slide(prs, "{{participant.full_name}}", notes="repeat:module")
    warnings = []
    _render(prs, DATA, warnings)
    assert any("participant.full_name" in w for w in warnings)


def test_nested_repeat_is_not_supported():
    """Slides are flat, so a list nested INSIDE a repeat item cannot drive its
    own repeat. This is why n8n must flatten modules to one row per
    (participant, module) pair instead of nesting them under participants."""
    prs = Presentation()
    _add_slide(prs, "{{module.title}}", notes="repeat:module")
    data = {"participant": [{"full_name": "Piotr", "module": [{"title": "BE"}]}]}
    warnings = []
    out = _render(prs, data, warnings)
    assert len(list(out.slides)) == 0
    assert any("repeat:module" in w and "dropped" in w for w in warnings)
