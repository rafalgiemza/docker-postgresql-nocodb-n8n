"""Kontrakt danych oferty po migracji do schematu v4
(docs/archive/fable/nocodb_crm_schema_v4.md, plan: docs/archive/fable/nocodb_crm_offers_v4.md).

Następca `test_schema_v3_contract.py`. Rendererowi NIE trzeba zmieniać ani linii
kodu; ten plik przypina kształt `data`, który musi wysyłać W1 po zmianie schematu
(lead -> szansa sprzedaży, participants -> persons), oraz to, co dziś W1 pomija.
Testy z prawdziwym szablonem (`init-data/source/ofer-template-1.0.pptx`, 32 MB,
poza gitem) są pomijane, gdy pliku nie ma.

Run:  pytest file-renderer-service/test_schema_v4_contract.py -v
"""
import io
import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from renderer import render_pptx

REAL_TEMPLATE = Path(__file__).resolve().parent.parent / "init-data" / "source" / "ofer-template-1.0.pptx"
needs_real_template = pytest.mark.skipif(
    not REAL_TEMPLATE.exists(), reason="prawdziwy szablon oferty nie jest w repo (init-data/source/ jest w .gitignore)")


def _add_slide(prs, text, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.paragraphs[0].add_run().text = text
    if notes is not None:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _render_synthetic(prs, data, warnings):
    buf = io.BytesIO()
    prs.save(buf)
    return Presentation(io.BytesIO(render_pptx(buf.getvalue(), data, warnings)))


def _text(slide):
    return slide.shapes[0].text_frame.text


# --------------------------------------------------------------------------
# Kształt `data`, który składa W1 v13 (docs/archive/fable/W1v13_generate_offer.json,
# węzeł "Assemble offer data"):
#   key                 <- offers.cover_style
#   lead.contact_name   <- szansa -> lead_name (fallback: kupujący, sale_opportunities.person)
#   assessment, a       <- najnowsza ocena pierwszej osoby (needs_summary tylko ai_accepted)
#   recommendation      <- rekomendacja pierwszej osoby (tylko ai_accepted)
#   participants[]      <- persons z offers.recommendations.person
#   packages[]          <- jeden wpis na slajd uzasadnienia wybranego core
#   others[]            <- dodatki (selected_package_variants_others), godziny sumowane
#   testimonial(s)      <- offers.selected_testimonials
#   logo_wall           <- offers.logo_wall_style
# Ceny wariantów (<Wariant>.price) NIE są wysyłane - następna iteracja (nowy szablon).
# --------------------------------------------------------------------------
DATA_V4 = {
    "key": "boy",
    "lead": {
        "contact_name": "ACME sp. z o.o."
    },
    "offer": {
        "date": "01.08.2026",
        "hours": 60,
        "price": 20400
    },
    "assessment": {
        "needs_summary": "demo",
        "position": "Developer"
    },
    "a": {
        "o": "B2.4",
        "r": "B2",
        "a": "B2",
        "f": "B1",
        "c": "B2"
    },
    "recommendation": {
        "headline": "BE + Skills",
        "rationale": "bo demo"
    },
    "participants": [
        {
            "full_name": "Piotr",
            "position": "Dev"
        },
        {
            "full_name": "Anna",
            "position": "PO"
        }
    ],
    "packages": [
        {
            "name": "Business English",
            "hours": 40,
            "title": "Cel 1",
            "description": "opis 1"
        },
        {
            "name": "Business English",
            "hours": 40,
            "title": "Cel 2",
            "description": "opis 2"
        },
        {
            "name": "Skills",
            "hours": 20,
            "title": "Skills",
            "description": ""
        }
    ],
    "others": [
        {
            "key": "facilitating",
            "name": "Facilitating",
            "description": "d",
            "hours": 15
        }
    ],
    "testimonials": [
        {
            "content": "Super",
            "name": "Jan K.",
            "position": "CTO"
        },
        {
            "content": "Ok",
            "name": "Anna N.",
            "position": "CEO"
        }
    ],
    "logo_wall": [
        {
            "key": "it"
        }
    ],
    "testimonial": {
        "content": "Super",
        "name": "Jan K.",
        "position": "CTO"
    }
}
TESTIMONIALS = DATA_V4["testimonials"]


def test_participants_and_packages_are_repeat_lists():
    """W1 v13: `repeat:participants` i `repeat:packages` (jedna oferta dla wielu
    osób, płaci lead, który nie musi być uczestnikiem). Klucze elementu są
    podniesione, nazwa listy działa też jako prefiks."""
    prs = Presentation()
    _add_slide(prs, "{{participants.full_name}} / {{position}}", notes="repeat:participants")
    _add_slide(prs, "{{packages.name}}: {{title}} ({{hours}}h)", notes="repeat:packages")
    warnings = []
    out = _render_synthetic(prs, {"participants": DATA_V4["participants"], "packages": DATA_V4["packages"]}, warnings)
    assert [_text(s) for s in out.slides] == [
        "Piotr / Dev", "Anna / PO",
        "Business English: Cel 1 (40h)", "Business English: Cel 2 (40h)", "Skills: Skills (20h)"]
    assert warnings == []


def test_others_and_logo_wall_lists():
    prs = Presentation()
    _add_slide(prs, "{{others.name}} {{hours}}h", notes="repeat:others")
    _add_slide(prs, "{{logo_wall.key}}", notes="repeat:logo_wall")
    warnings = []
    out = _render_synthetic(prs, {"others": DATA_V4["others"], "logo_wall": DATA_V4["logo_wall"]}, warnings)
    assert [_text(s) for s in out.slides] == ["Facilitating 15h", "it"]
    assert warnings == []


def test_testimonials_list_and_single_object_serve_different_slides():
    """Stały slajd czyta `testimonial.*` z POJEDYNCZEGO obiektu; slajd z
    `repeat:testimonials` w każdej kopii czyta element listy pod prefiksem
    `testimonials.*` (nazwa znacznika) albo z kluczy podniesionych (`name`)."""
    prs = Presentation()
    _add_slide(prs, "{{testimonial.name}}")                                       # stały
    _add_slide(prs, "{{testimonials.name}} / {{testimonials.content}}", notes="repeat:testimonials")
    warnings = []
    out = _render_synthetic(prs, {"testimonial": TESTIMONIALS[0], "testimonials": TESTIMONIALS},
                            warnings)
    assert [_text(s) for s in out.slides] == [
        "Jan K.", "Jan K. / Super", "Anna N. / Ok"]
    assert warnings == []


def test_singular_prefix_inside_plural_repeat_repeats_one_testimonial():
    """PUŁAPKA szablonu v1.0 (slajd 25): `repeat:testimonials` + `{{testimonial.x}}`.
    Prefiks liczby pojedynczej to osobny, GLOBALNY obiekt, więc każda kopia slajdu
    pokazuje tę samą referencję. Poprawka w szablonie: `{{testimonials.x}}`."""
    prs = Presentation()
    _add_slide(prs, "{{testimonial.name}}", notes="repeat:testimonials")
    warnings = []
    out = _render_synthetic(prs, {"testimonial": TESTIMONIALS[0], "testimonials": TESTIMONIALS},
                            warnings)
    assert [_text(s) for s in out.slides] == ["Jan K.", "Jan K."]


# -------------------------------------------------------------- prawdziwy szablon
EXPECTED_PLACEHOLDERS = {
    "lead.contact_name", "assessment.position", "assessment.needs_summary",
    "a.o", "a.r", "a.a", "a.f", "a.c",
    "offer.date", "offer.hours", "offer.price", "package.hours",
    "BusinessEnglish15h.price", "BusinessEnglish15hinpair.price",
    "skills1on1.price", "skillsgroup.price",
    "testimonial.content", "testimonial.name", "testimonial.position",
}


@needs_real_template
def test_real_template_demands_exactly_these_placeholders():
    """Zmiana szablonu (nowy placeholder) ma wywalić ten test i wymusić decyzję,
    skąd W1 weźmie dane."""
    warnings = []
    render_pptx(REAL_TEMPLATE.read_bytes(), {}, warnings)
    missing = {re.sub(r"^.*value: ", "", w) for w in warnings if "missing placeholder value" in w}
    assert missing == EXPECTED_PLACEHOLDERS


@needs_real_template
def test_real_template_with_w1_v13_data_has_only_known_gaps():
    """Szablon v1.0 vs dane W1 v13. Jedyne brakujące wartości to znane luki:
    `package.hours` (szablon czyta pojedynczy obiekt; docelowo `repeat:packages`
    + `{{packages.hours}}` - zmiana w szablonie) oraz ceny wariantów (kolejna
    iteracja, nowy szablon z cenami). Po poprawce szablonu zmień na `== []`."""
    warnings = []
    out = Presentation(io.BytesIO(render_pptx(REAL_TEMPLATE.read_bytes(), DATA_V4, warnings)))
    missing = {re.sub(r"^.*value: ", "", w) for w in warnings}
    assert missing == {"package.hours", "BusinessEnglish15h.price", "BusinessEnglish15hinpair.price",
                       "skills1on1.price", "skillsgroup.price"}
    assert all("missing placeholder value" in w for w in warnings)
    assert len(out.slides) == 29    # 28 w szablonie + 1 za drugi testimonial


def _all_text(shapes):
    for sh in shapes:
        if sh.shape_type == 6:                       # grupa
            yield from _all_text(sh.shapes)
        elif sh.has_text_frame:
            yield sh.text_frame.text


@needs_real_template
@pytest.mark.xfail(strict=True, reason="szablon v1.0: slajd 25 (repeat:testimonials) używa "
                   "{{testimonial.*}} zamiast {{testimonials.*}}, więc obie kopie pokazują "
                   "tę samą referencję; slajd 24 ma też literówkę {{testimonial.name} "
                   "(brakuje nawiasu). Test przejdzie po poprawce szablonu.")
def test_real_template_repeat_slides_show_distinct_testimonials():
    out = Presentation(io.BytesIO(render_pptx(REAL_TEMPLATE.read_bytes(), DATA_V4, [])))
    slides = list(out.slides)
    first, second = ("".join(_all_text(slides[i].shapes)) for i in (24, 25))
    assert "Jan K." in first and "Anna N." in second
    assert not any("{{" in t for sl in slides for t in _all_text(sl.shapes))
