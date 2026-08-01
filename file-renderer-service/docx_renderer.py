import io
from docx import Document

from placeholders import PLACEHOLDER, resolve


def render_paragraph(para, ctx, warnings):
    """Replace placeholders in paragraph text."""
    runs = para.runs
    if not runs:
        return

    full = "".join(r.text or "" for r in runs)
    if "{{" not in full:
        return

    sub = lambda s: PLACEHOLDER.sub(lambda m: resolve(m.group(1), ctx, warnings), s)

    # Fast path: every run is self-contained (no placeholder split across runs).
    if all(("{{" in (r.text or "")) == ("}}" in (r.text or "")) for r in runs) \
            and "{{" not in PLACEHOLDER.sub("", full):
        for r in runs:
            if r.text and "{{" in r.text:
                r.text = sub(r.text)
        return

    # Placeholder split across runs: collapse into first run.
    if runs:
        runs[0].text = sub(full)
        for r in runs[1:]:
            r.text = ""


def render_docx(template_bytes, data, warnings):
    """Render a DOCX template with placeholder substitution.

    Renders {{path.to.value}} placeholders in paragraphs, tables, headers, footers.
    """
    doc = Document(io.BytesIO(template_bytes))

    # Render main document paragraphs
    for para in doc.paragraphs:
        render_paragraph(para, data, warnings)

    # Render tables in main document
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    render_paragraph(para, data, warnings)

    # Render headers and footers
    for section in doc.sections:
        # Header
        for para in section.header.paragraphs:
            render_paragraph(para, data, warnings)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        render_paragraph(para, data, warnings)

        # Footer
        for para in section.footer.paragraphs:
            render_paragraph(para, data, warnings)
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        render_paragraph(para, data, warnings)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
