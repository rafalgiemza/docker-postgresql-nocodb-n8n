import copy, io, re
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.text.text import Font, _Hyperlink

import richtext
from richtext import Span
from placeholders import PLACEHOLDER, resolve

# LongText (rich text/markdown) fields from NocoDB - bold/italic/underline/
# strike/code/links, see richtext.py for the exact supported syntax and
# rationale. Headings/lists/quotes/code-blocks/images degrade to plain text
# for now (Phase 2/3, see .ai/offers/plan-september.md).
CODE_FONT = "Consolas"


def _apply_style(rPr_el, run, style):
    f = Font(rPr_el)
    f.bold = style.bold
    f.italic = style.italic
    f.underline = style.underline
    rPr_el.set("strike", "sngStrike" if style.strike else "noStrike")
    if style.code:
        f.name = CODE_FONT
    if style.link:
        _Hyperlink(rPr_el, run).address = style.link


def _apply_rich_text(run, spans):
    """Rewrite `run` in place as a richtext-aware run sequence (`spans` from
    richtext.segments()). New runs are deep-copies of the original run's XML
    (so they inherit its font/size/color) with only the toggled properties
    changed per segment; line breaks become `<a:br/>` siblings (soft break,
    same paragraph - a new `<a:p>` would repeat bullet/numbering and
    paragraph spacing)."""
    r_el = run._r
    p_el = r_el.getparent()
    run.text = ""
    anchor = r_el
    used_original = False
    for tok in spans:
        if not isinstance(tok, Span):
            br = p_el.add_br()
            anchor.addnext(br)
            anchor = br
            continue
        if not tok.text:
            continue
        if not used_original:
            run.text = tok.text
            _apply_style(r_el.get_or_add_rPr(), run, tok.style)
            anchor = r_el
            used_original = True
        else:
            new_el = copy.deepcopy(r_el)
            new_el.find(qn("a:t")).text = tok.text
            anchor.addnext(new_el)
            anchor = new_el
            _apply_style(new_el.get_or_add_rPr(), run, tok.style)


def render_paragraph(para, ctx, warnings):
    runs = para.runs
    full = "".join(r.text or "" for r in runs)
    if "{{" not in full:
        return
    sub = lambda s: PLACEHOLDER.sub(lambda m: resolve(m.group(1), ctx, warnings), s)
    # Fast path: every run is self-contained (no placeholder split across runs).
    if all(("{{" in (r.text or "")) == ("}}" in (r.text or "")) for r in runs) \
            and "{{" not in PLACEHOLDER.sub("", full):
        for r in runs:
            if r.text and "{{" in r.text:
                resolved = sub(r.text)
                spans = richtext.segments(resolved)
                if richtext.is_rich(spans):
                    _apply_rich_text(r, spans)
                else:
                    r.text = resolved
        return
    # Placeholder split across runs: collapse into the first run.
    # (Documented tradeoff: keep placeholders inside one styling run.)
    if runs:
        resolved = sub(full)
        spans = richtext.segments(resolved)
        if richtext.is_rich(spans):
            _apply_rich_text(runs[0], spans)
        else:
            runs[0].text = resolved
        for r in runs[1:]:
            r.text = ""


def render_shapes(shapes, ctx, warnings):
    for shp in shapes:
        if shp.shape_type == 6 and hasattr(shp, "shapes"):   # group
            render_shapes(shp.shapes, ctx, warnings)
            continue
        if getattr(shp, "has_text_frame", False):
            for para in shp.text_frame.paragraphs:
                render_paragraph(para, ctx, warnings)
        if getattr(shp, "has_table", False):
            for row in shp.table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        render_paragraph(para, ctx, warnings)


REL_ATTRS = [qn("r:embed"), qn("r:id"), qn("r:link")]


def duplicate_slide(prs, slide):
    """Copy a slide (shapes + image/media rels) and return the new slide.
    python-pptx has no native slide copy - see pptx skill notes."""
    new = prs.slides.add_slide(slide.slide_layout)
    for shp in list(new.shapes):                       # drop layout placeholders
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in slide.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            new_rid = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = new.part.rels.get_or_add(rel.reltype, rel.target_part)
        rid_map[rid] = new_rid
    for shp in slide.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in REL_ATTRS:
                old = node.get(attr)
                if old and old in rid_map:
                    node.set(attr, rid_map[old])
        new.shapes._spTree.append(el)
    return new


def move_slide_after(prs, slide, after):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    by_rid = {prs.part.rels[i.rId].target_part: i for i in ids}
    s_el, a_el = by_rid[slide.part], by_rid[after.part]
    lst.remove(s_el)
    a_el.addnext(s_el)


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for i in list(lst):
        if prs.part.rels[i.rId].target_part is slide.part:
            prs.part.drop_rel(i.rId)
            lst.remove(i)
            return


def slide_repeat_marker(slide):
    """`repeat:<name>` in a slide's SPEAKER NOTES - <name> is both the key
    looked up in `data` (must be a list) and the placeholder prefix usable on
    that slide (`{{<name>.field}}`, or bare `{{field}}` - see render_pptx).
    Fully generic: any name works, nothing hardcoded to a particular entity."""
    if not slide.has_notes_slide:
        return None
    txt = slide.notes_slide.notes_text_frame.text or ""
    m = re.search(r"repeat\s*:\s*([a-zA-Z_][a-zA-Z0-9_]*)", txt, re.I)
    return m.group(1) if m else None


def slide_render_if_marker(slide):
    """`render_if:<value>` in a slide's SPEAKER NOTES - the slide survives
    only if str(ctx.get("key")) == <value> (the literal, fixed "key" field;
    inside a repeat: block this is normally the item's own lifted "key",
    e.g. package_variants.key - see render_pptx). Same charset as
    repeat:<name>. Returns None if the marker isn't present/recognized, in
    which case the slide is unconditional (degrades safely, same as any
    other unrecognized marker text)."""
    if not slide.has_notes_slide:
        return None
    txt = slide.notes_slide.notes_text_frame.text or ""
    m = re.search(r"render_if\s*:\s*([a-zA-Z_][a-zA-Z0-9_]*)", txt, re.I)
    return m.group(1) if m else None


def render_if_ok(value, ctx, warnings):
    """True if `value` is None (slide has no render_if marker) or matches
    str(ctx.get("key")) exactly. Appends a warning and returns False on any
    mismatch, including a missing "key" - fail-closed, same convention as
    repeat:'s empty-list handling. Takes the already-resolved marker VALUE
    (not the slide) so callers control whether it's re-derived per slide or
    captured once and reused across a repeat group's duplicates - duplicated
    slides have no notes slide part (see duplicate_slide), so re-deriving it
    from a duplicated target would silently see "no marker" and stop
    filtering."""
    if value is None:
        return True
    if value == str(ctx.get("key")):
        return True
    warnings.append(f"render_if:{value} slide dropped - key={ctx.get('key')!r} did not match")
    return False


def render_pptx(template_bytes, data, warnings):
    """MUSI zrobic wszystkie duplikacje PRZED jakimkolwiek usunieciem slajdu w
    tym samym przebiegu. python-pptx liczy nazwe pliku nowego slajdu jako
    `len(sldIdLst) + 1` (PresentationPart._next_slide_partname) - bez
    sprawdzenia, czy taki plik juz istnieje. Jesli wczesniej w tym samym
    przebiegu cokolwiek zostalo usuniete (np. niepasujacy wariant
    render_if), liczba slajdow sie zmniejsza, a kolejna duplikacja (np. dla
    repeat:) dostaje numer, ktory juz nalezy do INNEGO, wciaz zywego slajdu
    dalej w prezentacji - dwa rozne slajdy ladu ja w archiwum pod tym samym
    "slideN.xml", co PowerPoint naprawia, usuwajac jeden z nich (ZWERYFIKOWANE
    NA ZYWO: reprodukowane na prawdziwym szablonie, patrz historia zmian -
    "duplicate slide29.xml/slide30.xml"). Dlatego ta funkcja NAJPIERW w
    calosci przechodzi original slajdy i wykonuje WSZYSTKIE duplikacje
    (liczba slajdow tylko rosnie), zbierajac decyzje "usun"/"renderuj", a
    USUWANIE i RENDEROWANIE odklada na koniec, gdy zadna kolejna duplikacja
    juz nie nastapi.
    """
    prs = Presentation(io.BytesIO(template_bytes))
    to_delete = []
    to_render = []  # (slide, ctx) par - shapes renderowane dopiero po fazie usuwania
    for slide in list(prs.slides):
        marker = slide_repeat_marker(slide)
        # Captured once from the PRISTINE slide's own notes, same reason
        # `marker` is: duplicate_slide() does not copy the notesSlide
        # relationship, so a duplicated target has no notes of its own to
        # re-derive this from later.
        render_if_value = slide_render_if_marker(slide)
        if not marker:
            if not render_if_ok(render_if_value, data, warnings):
                to_delete.append(slide)
                continue
            to_render.append((slide, data))
            continue
        items = data.get(marker) or []
        if not items:
            warnings.append(f"repeat:{marker} slide dropped - no items")
            to_delete.append(slide)
            continue
        # Clone all copies from the PRISTINE template slide FIRST (before any
        # render mutates it), then render each with its own item context.
        targets = [slide]
        anchor = slide
        for _ in range(len(items) - 1):
            dup = duplicate_slide(prs, slide)
            move_slide_after(prs, dup, anchor)
            anchor = dup
            targets.append(dup)
        for target, item in zip(targets, items):
            # The item is exposed BOTH under the marker name
            # ({{participant.position}}) and with its own keys lifted to the
            # top level ({{position}}, and thus {{a.o}} when the item carries
            # an "a" object). Lifted keys shadow same-named keys in `data` for
            # this slide only - documented tradeoff; the prefixed form always
            # stays available and takes precedence for the marker name itself.
            ctx = {**data}
            if isinstance(item, dict):
                ctx.update(item)
            ctx[marker] = item
            if not render_if_ok(render_if_value, ctx, warnings):
                to_delete.append(target)
                continue
            to_render.append((target, ctx))
    for slide in to_delete:
        delete_slide(prs, slide)
    for target, ctx in to_render:
        render_shapes(target.shapes, ctx, warnings)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


