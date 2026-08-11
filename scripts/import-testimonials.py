#!/usr/bin/env python3
r"""Etap 1/2 importu referencji klienta: wypełnia `testimonials` z
`init-data/source/testimonials.xlsx` i linkuje firmę po nazwie, jeśli
istnieje pasujący rekord w `companies`.

Uruchamiany PO zasileniu bazy starymi leadami/firmami (`make init-data` ->
seed-service) - dopasowanie firmy po nazwie ma sens dopiero, gdy `companies`
nie jest puste. Etap 2 (dociągnięcie slajdów z pptx) to osobny skrypt,
`attach-testimonial-slides.py`, uruchamiany po tym.

Nagłówki xlsx (klient, 2026-08-11) i mapowanie -> patrz plan
`/Users/rafalgiemza/.claude/plans/virtual-prancing-lightning.md`. Skrót:
  Opinia -> content, Imię i nazwisko -> client_name, Rola/zawód -> position,
  Firma -> WYŁĄCZNIE do dopasowania companies (nie zapisywane wprost - branża/
  wielkość firmy żyją teraz na companies, nie duplikujemy), Branża/Wielkość
  firmy -> pomijane (jw.), "Czy możemy używać..." -> usage_consent
  (yes/no/partially, heurystyka niżej) + usage_limitations (pełny surowy
  tekst), Uwagi -> notes, buyer persona / "Do czego się odnosi" /
  Tłumaczenie -> wprost (wolny tekst, bez ustalonej listy).

"Czy możemy używać..." NIE jest czystym TAK/NIE w realnych danych (sprawdzone
na dostarczonym pliku) - bywa "-", "???", puste, albo pełne zdanie z
zastrzeżeniem ("jedynie do ofert", z linkiem, itp.). Klasyfikacja jest więc
heurystyką, nie ścisłym parserem - `usage_limitations` zawsze trzyma surowy
tekst 1:1, więc nic nie ginie nawet gdy heurystyka się pomyli.

Idempotentny: dedupe po kluczu naturalnym fold(client_name)+fold(content) -
bez osobnego pola legacy_id. Wiersz z kluczem już istniejącym w bazie (albo
zdublowany w obrębie samego xlsx) jest pomijany.

Wymaga w środowisku (patrz .env.example): NC_API_TOKEN, NC_CRM_BASE_ID.
Opcjonalnie: NC_LOCAL_URL (default http://localhost:8081).

Usage (przez kontener testimonials-import - patrz jego README, PEP 668 na
hoście robi z gołego `pip install` problem):
  make init-data
  # albo bezpośrednio:
  ./scripts/import-testimonials.sh --dry-run   # ZAWSZE najpierw
  ./scripts/import-testimonials.sh [--xlsx PATH]

  Ten plik da się też odpalić bezpośrednio (python3 scripts/import-testimonials.py),
  jeśli masz gdzieś już zainstalowane init-data/requirements.txt (np. venv).
"""
import argparse
import os
import sys
import unicodedata
from collections import defaultdict

import openpyxl
import requests

URL = os.environ.get("NC_LOCAL_URL", "http://localhost:8081").rstrip("/")
TOKEN = os.environ.get("NC_API_TOKEN")
BASE_ID = os.environ.get("NC_CRM_BASE_ID")
DEFAULT_XLSX = "init-data/source/testimonials.xlsx"

if not TOKEN or not BASE_ID:
    sys.exit("Brak NC_API_TOKEN / NC_CRM_BASE_ID w środowisku - patrz .env.example.")

S = requests.Session()
S.headers.update({"xc-token": TOKEN, "Content-Type": "application/json"})


def api(method, path, **kw):
    try:
        r = S.request(method, f"{URL}{path}", timeout=30, **kw)
    except requests.RequestException as e:
        sys.exit(f"Nie moge sie polaczyc z NocoDB ({URL}): {type(e).__name__}. "
                 f"Sprawdz NC_LOCAL_URL / czy kontener stoi.")
    if not r.ok:
        sys.exit(f"NocoDB API error {r.status_code} on {method} {path}: {r.text[:500]}")
    return r.json() if r.text else {}


# ----------------------------------------------------------------- helpers
def fold(s):
    """Accent/case-insensitive compare key. Strips leading/trailing
    punctuation too - real dane maja literowki jak 'Grabowska-' (sprawdzone
    2026-08-11 na dostarczonym xlsx), ktore inaczej lamalyby dopasowanie."""
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore") \
        .decode().lower().strip()
    return s.strip(" -,.")


def clean_header(h):
    return " ".join(str(h or "").split())


# ----------------------------------------------------------------- meta
def resolve_meta():
    """Zwraca (testimonials_id, companies_id, company_link_field_id,
    usage_consent_options)."""
    tables = {t["title"].strip().lower(): t["id"]
              for t in api("GET", f"/api/v2/meta/bases/{BASE_ID}/tables").get("list", [])}
    missing = [t for t in ("testimonials", "companies") if t not in tables]
    if missing:
        sys.exit(f"Brak tabel (uruchom najpierw `make init-schema`): {missing}")
    tid, cid = tables["testimonials"], tables["companies"]

    company_field_id = None
    usage_consent_options = []
    for col in api("GET", f"/api/v2/meta/tables/{tid}").get("columns", []):
        uidt = col.get("uidt")
        if uidt in ("Links", "LinkToAnotherRecord"):
            # Dopasowanie po fk_related_model_id, NIE po tytule - NocoDB sam
            # nadaje nazwę polu zwrotnemu przy relacji hm (companies -> testimonials),
            # nie da się jej narzucić z init-schema.py (patrz jego nagłówek).
            related = (col.get("colOptions") or {}).get("fk_related_model_id")
            if related == cid:
                company_field_id = col["id"]
        elif col["title"].strip().lower() == "usage_consent":
            usage_consent_options = [o["title"] for o in
                                     (col.get("colOptions") or {}).get("options", [])]
    if not company_field_id:
        sys.exit("Nie znalazłem na `testimonials` pola linkującego do `companies` - "
                 "uruchom `make init-schema` (tworzy relację companies->testimonials).")
    if not usage_consent_options:
        sys.exit("Pole `testimonials.usage_consent` nie istnieje albo nie ma opcji - "
                 "uruchom `make init-schema` (ewentualnie po resecie appdata, patrz "
                 "nagłówek scripts/init-schema.py pkt 5: nie dokleja pól do już "
                 "istniejącej tabeli).")
    return tid, cid, company_field_id, set(usage_consent_options)


def existing_testimonial_keys(tid):
    rows = api("GET", f"/api/v2/tables/{tid}/records"
                       "?fields=client_name,content&limit=1000").get("list", [])
    return {(fold(r.get("client_name")), fold(r.get("content"))) for r in rows}


def existing_company_map(cid):
    rows = api("GET", f"/api/v2/tables/{cid}/records?fields=name&limit=1000").get("list", [])
    out = {}
    for r in rows:
        k = fold(r.get("name"))
        if k:
            out[k] = r["Id"]
    return out


def link_company(tid, field_id, testimonial_id, company_id):
    api("POST", f"/api/v2/tables/{tid}/links/{field_id}/records/{testimonial_id}",
        json=[{"Id": company_id}])


# ----------------------------------------------------------------- xlsx
HEADERS = {
    "content": "Opinia",
    "in_pptx_flag": "czy jest w pptx",
    "client_name": "Imię i nazwisko",
    "position": "Rola/zawód",
    "company": "Firma",
    "usage_raw": "Czy możemy używać (SM, www, ofertowanie, zdjęcie z nazwiskiem, "
                 "nazwa firmy)",
    "notes": "Uwagi",
    "buyer_persona": "Do której buyer persony najbardziej trafia",
    "refers_to": "Do czego się odnosi",
    "translation": "Tłumaczenie",
}


def load_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    raw_rows = list(ws.iter_rows(values_only=True))
    headers = [clean_header(h) for h in raw_rows[0]]
    missing = [h for h in HEADERS.values() if h not in headers]
    if missing:
        sys.exit(f"Brakuje oczekiwanych nagłówków w {path}: {missing}\n"
                 f"Nagłówki w pliku: {headers}")
    idx = {key: headers.index(h) for key, h in HEADERS.items()}
    rows = []
    for r in raw_rows[1:]:
        if not any(v not in (None, "") for v in r):
            continue
        rows.append({key: r[i] if i < len(r) else None for key, i in idx.items()})
    return rows


# ----------------------------------------------------------------- classification
TRUE_MARKERS = {"tak", "x", "1", "true", "yes", "✓"}


def as_bool(v):
    s = fold(v)
    if not s:
        return False
    return s in TRUE_MARKERS


def classify_usage(raw):
    """Surowa komórka -> ('yes'|'no'|'partially'|None, pełny surowy tekst)."""
    text = str(raw or "").strip()
    if not text or text == "???":
        return None, text
    f = fold(text)
    if f.startswith("tak"):
        return "yes", text
    if f == "-" or f == "nie" or f.startswith("nie"):
        return "no", text
    return "partially", text


# ----------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    args = ap.parse_args()

    print(f"NocoDB: {URL}, base: {BASE_ID}{' [DRY RUN]' if args.dry_run else ''}")
    if not os.path.exists(args.xlsx):
        sys.exit(f"Nie znaleziono pliku: {args.xlsx}")

    tid, cid, company_field_id, usage_options = resolve_meta()
    rows = load_rows(args.xlsx)
    print(f"Wiersze w xlsx: {len(rows)}")

    existing_keys = existing_testimonial_keys(tid)
    companies = existing_company_map(cid)
    print(f"Firmy w bazie do dopasowania: {len(companies)}")

    stats = defaultdict(int)
    unmatched_companies = defaultdict(int)
    usage_dist = defaultdict(int)
    usage_samples = defaultdict(list)
    bad_pptx_flags = set()
    seen_this_run = set()

    for row in rows:
        stats["rows_read"] += 1
        client_name = str(row["client_name"] or "").strip()
        content = str(row["content"] or "").strip()
        key = (fold(client_name), fold(content))
        if key in existing_keys or key in seen_this_run:
            stats["skipped_duplicate"] += 1
            continue
        seen_this_run.add(key)

        company_raw = clean_header(row["company"])
        company_id = companies.get(fold(company_raw)) if company_raw else None
        if company_raw and not company_id:
            unmatched_companies[company_raw] += 1

        usage_consent, usage_limitations = classify_usage(row["usage_raw"])
        usage_dist[usage_consent or "(puste)"] += 1
        if usage_consent in (None, "partially") and len(usage_samples[usage_consent or "(puste)"]) < 5:
            usage_samples[usage_consent or "(puste)"].append(usage_limitations)
        if usage_consent and usage_consent not in usage_options:
            sys.exit(f"usage_consent='{usage_consent}' nie jest realną opcją pola "
                     f"({usage_options}) - schemat i skrypt się rozjechały.")

        pptx_flag_raw = row["in_pptx_flag"]
        in_source_pptx = as_bool(pptx_flag_raw)
        if pptx_flag_raw not in (None, "") and fold(pptx_flag_raw) not in TRUE_MARKERS \
                and fold(pptx_flag_raw) not in {"nie", "n", "0", "false"}:
            bad_pptx_flags.add(str(pptx_flag_raw))

        title = f"{client_name} – {company_raw}" if company_raw else client_name

        record = {
            "title": title,
            "client_name": client_name,
            "position": clean_header(row["position"]),
            "type": "testimonial",
            "content": content,
            "translation": clean_header(row["translation"]),
            "buyer_persona": clean_header(row["buyer_persona"]),
            "refers_to": clean_header(row["refers_to"]),
            "usage_limitations": usage_limitations,
            "in_source_pptx": in_source_pptx,
            "notes": clean_header(row["notes"]),
            "active": True,
        }
        if usage_consent:
            record["usage_consent"] = usage_consent
        # `in_source_pptx` (bool) i `active` (zawsze True) przetrwają ten filtr -
        # False/True nigdy nie są równe None ani "".
        record = {k: v for k, v in record.items() if v not in (None, "")}

        print(f"+  {title!r}" + (f" -> firma: {company_raw}" if company_id
              else f" -> firma NIEDOPASOWANA: {company_raw!r}" if company_raw else ""))
        stats["created"] += 1
        if company_id:
            stats["linked_to_company"] += 1
        if args.dry_run:
            continue

        rec_id = api("POST", f"/api/v2/tables/{tid}/records", json=record)
        rec_id = rec_id.get("Id") or rec_id.get("id")
        if company_id:
            link_company(tid, company_field_id, rec_id, company_id)

    print("\n===== REPORT =====")
    for k, v in sorted(stats.items()):
        print(f"{k}: {v}")
    print(f"\nusage_consent rozkład: {dict(usage_dist)}")
    for k, samples in usage_samples.items():
        print(f"  próbki '{k}': {samples}")
    if unmatched_companies:
        print(f"\nNIEDOPASOWANE firmy ({len(unmatched_companies)}) - dodaj do "
              f"`companies` albo popraw nazwę w xlsx i uruchom ponownie:")
        for name, n in sorted(unmatched_companies.items()):
            print(f"  {name!r} (x{n})")
    if bad_pptx_flags:
        print(f"\nNierozpoznane wartości 'czy jest w pptx' (potraktowane jako False): "
              f"{sorted(bad_pptx_flags)}")
    if args.dry_run:
        print("\n(dry run - nic nie zapisano)")


if __name__ == "__main__":
    main()
