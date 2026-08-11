#!/usr/bin/env python3
r"""Wypełnia/aktualizuje `pricing` z `init-data/source/cennik.xlsx`.

Plik dostarczony przez klientkę (2026-08-11) ma puste kolumny cenowe -
świadomie, żeby nie wysyłać realnych stawek poza firmę przy pierwszym
udostępnieniu. Struktura (segment/hours/tryb/...) jest już finalna, tylko
liczby dojdą później - dlatego ten skrypt jest UPSERTEM, nie dedupe-i-pomiń
jak import-testimonials.py: klucz naturalny to wymiary STRUKTURALNE
(product+customer_segment+training_group_size+hours+course_type), nie ceny.
Gdy na VPS podmienisz ten plik na wersję z prawdziwymi liczbami i odpalisz
skrypt ponownie, istniejące wiersze dostają PATCH z nowymi cenami zamiast
being pominięte jako "już istnieją".

Nagłówki xlsx -> pola `pricing` (patrz też opisy pól w scripts/init-schema.py,
tabela "pricing"):
  Liczba kursów -> customer_segment (wartości 1:1: "A) <10" itd.)
  Wielkość pakietu (h) -> hours
  Dodatkowe godziny za kontynuację -> continuation_bonus_hours
  Kontynuacja czy nowy kurs? -> course_type (kontynuacja/nowy_kurs)
  Dodatkowa h za płatność z góry -> upfront_payment_bonus (tak/nie -> bool)
  Rodzaj szkolenia -> training_group_size (indywidualne/w parach/grupowe ->
    1-1/w_parach/grupa, TRAINING_MODE_MAP niżej)
  Standardowa cena za pakiet -> total_price
  Po rabacie -> discounted_price
  Wielkość pakietu z bonusowymi h -> hours_with_bonus
  Cena za godzinę -> price_per_hour
  Liczba osób w grupie -> group_size
  Cena za godzinę za osobę -> price_per_hour_per_person
  Cena za h z bonusami -> price_per_hour_with_bonus
  Cena za h za osobę z bonusami -> price_per_hour_per_person_with_bonus
  Liczba pakietów / Suma za całość zakupu -> POMIJANE (pola kalkulatora
    zamówienia w arkuszu, nie stawki referencyjne - patrz init-schema.py).

`product` NIE ma kolumny w arkuszu - cały plik to jedna linia produktowa
("BUSINESS ENGLISH/ENGLISH FOR IT" ze slajdu cennikowego szablonu oferty,
gdzie te dwa tory kosztują tyle samo). Domyślnie "standard" (`--product`
pozwala zaimportować analogicznie ukształtowany arkusz dla innej wartości
z VARIANT, np. intensive_workshop, gdy taki powstanie).

Osobliwość źródłowego pliku: kolumna "Dodatkowe godziny za kontynuację" i
czasem "Wielkość pakietu z bonusowymi h" mają wartość 1.5 zapisaną jako
DATA (Excel autoformat, np. 2026-01-05 = "1.5" bo miesiąc.dzień) zamiast
liczby - parse_number() niżej to odkręca.

Idempotentny w warstwie STRUKTURY (nie dubluje wierszy przy ponownym
uruchomieniu), ale NIE idempotentny w warstwie cen - kolejne uruchomienie
NADPISUJE total_price/discounted_price/... wartościami z pliku. To celowe
(patrz wyżej), ale oznacza, że ręczne poprawki cen wprost w NocoDB UI
przetrwają tylko do następnego importu tego samego wiersza.

Wymaga w środowisku (patrz .env.example): NC_API_TOKEN, NC_CRM_BASE_ID.
Opcjonalnie: NC_LOCAL_URL (default http://localhost:8081).

Usage (przez kontener testimonials-import - patrz jego README, PEP 668 na
hoście robi z gołego `pip install` problem):
  ./scripts/import-pricing.sh --dry-run   # ZAWSZE najpierw
  ./scripts/import-pricing.sh [--xlsx PATH] [--product standard]
"""
import argparse
import datetime
import os
import sys
import unicodedata
from collections import defaultdict

import openpyxl
import requests

URL = os.environ.get("NC_LOCAL_URL", "http://localhost:8081").rstrip("/")
TOKEN = os.environ.get("NC_API_TOKEN")
BASE_ID = os.environ.get("NC_CRM_BASE_ID")
DEFAULT_XLSX = "init-data/source/cennik.xlsx"

if not TOKEN or not BASE_ID:
    sys.exit("Brak NC_API_TOKEN / NC_CRM_BASE_ID w środowisku - patrz .env.example.")

S = requests.Session()
S.headers.update({"xc-token": TOKEN, "Content-Type": "application/json"})


def api(method, path, **kw):
    try:
        r = S.request(method, f"{URL}{path}", timeout=30, **kw)
    except requests.RequestException as e:
        sys.exit(f"Nie moge sie polaczyc z NocoDB ({URL}): {type(e).__name__}. "
                 f"Sprawdz NC_LOCAL_URL / czy kontener stoi.")
    if not r.ok:
        sys.exit(f"NocoDB API error {r.status_code} on {method} {path}: {r.text[:500]}")
    return r.json() if r.text else {}


# ----------------------------------------------------------------- helpers
def fold(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore") \
        .decode().lower().strip()
    return s.strip(" -,.")


def clean_header(h):
    return " ".join(str(h or "").split())


TRAINING_MODE_MAP = {"indywidualne": "1-1", "w parach": "w_parach", "grupowe": "grupa"}
COURSE_TYPE_MAP = {"kontynuacja": "kontynuacja", "nowy kurs": "nowy_kurs",
                    "nowy_kurs": "nowy_kurs"}
TRUE_MARKERS = {"tak", "x", "1", "true", "yes"}
FALSE_MARKERS = {"nie", "n", "0", "false", ""}


def as_bool(v):
    f = fold(v)
    if f in TRUE_MARKERS:
        return True
    if f in FALSE_MARKERS:
        return False
    return None


def parse_number(v):
    """None dla pustych komórek (cennik.xlsx dziś ma je puste - świadomie,
    patrz nagłówek modułu), float dla liczb/tekstu-z-liczbą, i odkręcenie
    Excelowego autoformatu daty dla wartości typu 1.5 (patrz nagłówek)."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, datetime.datetime):
        return float(f"{v.month}.{v.day}")
    s = str(v).strip().replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


# ----------------------------------------------------------------- meta
def resolve_meta():
    """Zwraca (pricing_table_id, {field_title: set(opcje)} dla SingleSelect)."""
    tables = {t["title"].strip().lower(): t["id"]
              for t in api("GET", f"/api/v2/meta/bases/{BASE_ID}/tables").get("list", [])}
    if "pricing" not in tables:
        sys.exit("Brak tabeli `pricing` (uruchom najpierw `make init-schema`).")
    tid = tables["pricing"]

    options = {}
    for col in api("GET", f"/api/v2/meta/tables/{tid}").get("columns", []):
        if col.get("uidt") == "SingleSelect":
            options[col["title"].strip().lower()] = {
                o["title"] for o in (col.get("colOptions") or {}).get("options", [])}
    needed = ("product", "customer_segment", "training_group_size", "course_type")
    missing = [f for f in needed if not options.get(f)]
    if missing:
        sys.exit(f"Pola SingleSelect bez opcji na `pricing`: {missing} - uruchom "
                 f"`make init-schema` (nie dokleja pól/opcji do już istniejącej "
                 f"tabeli, patrz nagłówek scripts/init-schema.py pkt 5).")
    return tid, options


def existing_rows(tid):
    """klucz strukturalny -> Id rekordu, do UPSERT (patrz nagłówek modułu)."""
    rows = api("GET", f"/api/v2/tables/{tid}/records?limit=1000").get("list", [])
    out = {}
    for r in rows:
        key = (r.get("product"), r.get("customer_segment"),
               r.get("training_group_size"), r.get("hours"), r.get("course_type"))
        out[key] = r["Id"]
    return out


# ----------------------------------------------------------------- xlsx
HEADER_MARKER = "Liczba kursów"
HEADERS = {
    "customer_segment": "Liczba kursów",
    "hours": "Wielkość pakietu (h)",
    "continuation_bonus_hours": "Dodatkowe godziny za kontynuację",
    "course_type": "Kontynuacja czy nowy kurs?",
    "upfront_payment_bonus": "Dodatkowa h za płatność z góry",
    "training_group_size": "Rodzaj szkolenia",
    "total_price": "Standardowa cena za pakiet",
    "discounted_price": "Po rabacie",
    "hours_with_bonus": "Wielkość pakietu z bonusowymi h",
    "price_per_hour": "Cena za godzinę",
    "group_size": "Liczba osób w grupie",
    "price_per_hour_per_person": "Cena za godzinę za osobę",
    "price_per_hour_with_bonus": "Cena za h z bonusami",
    "price_per_hour_per_person_with_bonus": "Cena za h za osobę z bonusami",
}


def load_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    raw_rows = list(ws.iter_rows(values_only=True))

    header_i = next((i for i, r in enumerate(raw_rows)
                     if r and clean_header(r[0]) == HEADER_MARKER), None)
    if header_i is None:
        sys.exit(f"Nie znalazłem wiersza nagłówka (kolumna A = {HEADER_MARKER!r}) w {path}.")
    headers = [clean_header(h) for h in raw_rows[header_i]]
    missing = [h for h in HEADERS.values() if h not in headers]
    if missing:
        sys.exit(f"Brakuje oczekiwanych nagłówków w {path}: {missing}\n"
                 f"Nagłówki w pliku: {headers}")
    idx = {key: headers.index(h) for key, h in HEADERS.items()}

    rows = []
    for r in raw_rows[header_i + 1:]:
        if not r or not any(v not in (None, "") for v in r):
            break  # pierwszy pusty wiersz = koniec danych (dalej "Suma łączna")
        rows.append({key: r[i] if i < len(r) else None for key, i in idx.items()})
    return rows


# ----------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--product", default="standard",
                    help="wartość pricing.product dla całego arkusza (default: standard)")
    args = ap.parse_args()

    print(f"NocoDB: {URL}, base: {BASE_ID}{' [DRY RUN]' if args.dry_run else ''}")
    if not os.path.exists(args.xlsx):
        sys.exit(f"Nie znaleziono pliku: {args.xlsx}")

    tid, options = resolve_meta()
    if args.product not in options["product"]:
        sys.exit(f"--product={args.product!r} nie jest opcją pricing.product "
                 f"({options['product']}).")

    raw_rows = load_rows(args.xlsx)
    print(f"Wiersze w xlsx: {len(raw_rows)}")
    existing = existing_rows(tid)

    stats = defaultdict(int)
    seen_this_run = set()
    unrecognized = defaultdict(set)

    for raw in raw_rows:
        stats["rows_read"] += 1

        segment_raw = clean_header(raw["customer_segment"])
        if segment_raw not in options["customer_segment"]:
            unrecognized["customer_segment"].add(segment_raw)
            stats["skipped_bad_row"] += 1
            continue

        mode_raw = fold(raw["training_group_size"])
        mode = TRAINING_MODE_MAP.get(mode_raw)
        if not mode or mode not in options["training_group_size"]:
            unrecognized["training_group_size"].add(str(raw["training_group_size"]))
            stats["skipped_bad_row"] += 1
            continue

        course_type = COURSE_TYPE_MAP.get(fold(raw["course_type"]))
        if not course_type or course_type not in options["course_type"]:
            unrecognized["course_type"].add(str(raw["course_type"]))
            stats["skipped_bad_row"] += 1
            continue

        hours = parse_number(raw["hours"])
        if hours is None:
            stats["skipped_bad_row"] += 1
            continue

        key = (args.product, segment_raw, mode, hours, course_type)
        if key in seen_this_run:
            stats["skipped_duplicate_in_file"] += 1
            continue
        seen_this_run.add(key)

        record = {
            "product": args.product,
            "customer_segment": segment_raw,
            "training_group_size": mode,
            "hours": hours,
            "course_type": course_type,
            "continuation_bonus_hours": parse_number(raw["continuation_bonus_hours"]),
            "upfront_payment_bonus": as_bool(raw["upfront_payment_bonus"]),
            "total_price": parse_number(raw["total_price"]),
            "discounted_price": parse_number(raw["discounted_price"]),
            "hours_with_bonus": parse_number(raw["hours_with_bonus"]),
            "price_per_hour": parse_number(raw["price_per_hour"]),
            "group_size": parse_number(raw["group_size"]),
            "price_per_hour_per_person": parse_number(raw["price_per_hour_per_person"]),
            "price_per_hour_with_bonus": parse_number(raw["price_per_hour_with_bonus"]),
            "price_per_hour_per_person_with_bonus":
                parse_number(raw["price_per_hour_per_person_with_bonus"]),
        }
        record = {k: v for k, v in record.items() if v is not None}

        existing_id = existing.get(key)
        label = f"{args.product}/{segment_raw}/{mode}/{hours}h/{course_type}"
        if existing_id:
            print(f"~  {label} (Id={existing_id}, update)")
            stats["updated"] += 1
            if not args.dry_run:
                # PATCH chce listy rekordów, nie gołego dicta - zweryfikowane
                # na żywo w attach-testimonial-slides.py (update_record()).
                api("PATCH", f"/api/v2/tables/{tid}/records",
                    json=[{"Id": existing_id, **record}])
        else:
            print(f"+  {label} (nowy)")
            stats["created"] += 1
            if not args.dry_run:
                api("POST", f"/api/v2/tables/{tid}/records", json=record)

    print("\n===== REPORT =====")
    for k, v in sorted(stats.items()):
        print(f"{k}: {v}")
    if unrecognized:
        print("\nNIEROZPOZNANE wartości (wiersz pominięty) - sprawdź arkusz albo "
              "zmapuj w skrypcie:")
        for field, vals in unrecognized.items():
            print(f"  {field}: {sorted(vals)}")
    if args.dry_run:
        print("\n(dry run - nic nie zapisano)")


if __name__ == "__main__":
    main()
