import copy, io, re
from pptx import Presentation
from pptx.oxml.ns import qn

PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}")


def resolve(path, ctx, warnings):
    cur = ctx
    for part in path.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        else:
            warnings.append(f"missing placeholder value: {path}")
            return ""
    if cur is None:
        return ""
    if isinstance(cur, float) and cur == int(cur):
        cur = int(cur)
    return str(cur)


def render_paragraph(para, ctx, warnings):
    runs = para.runs
    full = "".join(r.text or "" for r in runs)
    if "{{" not in full:
        return
    sub = lambda s: PLACEHOLDER.sub(lambda m: resolve(m.group(1), ctx, warnings), s)
    # Fast path: every run is self-contained (no placeholder split across runs).
    if all(("{{" in (r.text or "")) == ("}}" in (r.text or "")) for r in runs) \
            and "{{" not in PLACEHOLDER.sub("", full):
        for r in runs:
            if r.text and "{{" in r.text:
                r.text = sub(r.text)
        return
    # Placeholder split across runs: collapse into the first run.
    # (Documented tradeoff: keep placeholders inside one styling run.)
    if runs:
        runs[0].text = sub(full)
        for r in runs[1:]:
            r.text = ""


def render_shapes(shapes, ctx, warnings):
    for shp in shapes:
        if shp.shape_type == 6 and hasattr(shp, "shapes"):   # group
            render_shapes(shp.shapes, ctx, warnings)
            continue
        if getattr(shp, "has_text_frame", False):
            for para in shp.text_frame.paragraphs:
                render_paragraph(para, ctx, warnings)
        if getattr(shp, "has_table", False):
            for row in shp.table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        render_paragraph(para, ctx, warnings)


REL_ATTRS = [qn("r:embed"), qn("r:id"), qn("r:link")]


def duplicate_slide(prs, slide):
    """Copy a slide (shapes + image/media rels) and return the new slide.
    python-pptx has no native slide copy - see pptx skill notes."""
    new = prs.slides.add_slide(slide.slide_layout)
    for shp in list(new.shapes):                       # drop layout placeholders
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in slide.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            new_rid = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = new.part.rels.get_or_add(rel.reltype, rel.target_part)
        rid_map[rid] = new_rid
    for shp in slide.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in REL_ATTRS:
                old = node.get(attr)
                if old and old in rid_map:
                    node.set(attr, rid_map[old])
        new.shapes._spTree.append(el)
    return new


def move_slide_after(prs, slide, after):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    by_rid = {prs.part.rels[i.rId].target_part: i for i in ids}
    s_el, a_el = by_rid[slide.part], by_rid[after.part]
    lst.remove(s_el)
    a_el.addnext(s_el)


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for i in list(lst):
        if prs.part.rels[i.rId].target_part is slide.part:
            prs.part.drop_rel(i.rId)
            lst.remove(i)
            return


def slide_repeat_marker(slide):
    if not slide.has_notes_slide:
        return None
    txt = slide.notes_slide.notes_text_frame.text or ""
    m = re.search(r"repeat\s*:\s*(participants|testimonials)", txt, re.I)
    return m.group(1).lower() if m else None


def render_pptx(template_bytes, data, warnings):
    prs = Presentation(io.BytesIO(template_bytes))
    key_map = {"participants": "participant", "testimonials": "testimonial"}
    for slide in list(prs.slides):
        marker = slide_repeat_marker(slide)
        if not marker:
            render_shapes(slide.shapes, data, warnings)
            continue
        items = data.get(marker) or []
        if not items:
            warnings.append(f"repeat:{marker} slide dropped - no items")
            delete_slide(prs, slide)
            continue
        # Clone all copies from the PRISTINE template slide FIRST (before any
        # render mutates it), then render each with its own item context.
        targets = [slide]
        anchor = slide
        for _ in range(len(items) - 1):
            dup = duplicate_slide(prs, slide)
            move_slide_after(prs, dup, anchor)
            anchor = dup
            targets.append(dup)
        for target, item in zip(targets, items):
            render_shapes(target.shapes, {**data, key_map[marker]: item}, warnings)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


