"""Offline unit tests for renderer.py — no NocoDB, no network.

Run:  pytest offer-service/test_renderer.py -v
"""
import io

from pptx import Presentation
from pptx.util import Inches

from renderer import render_pptx


def _new_prs():
    return Presentation()


def _add_slide(prs, runs_text, notes=None):
    """Blank-layout slide with one textbox; runs_text is a list of run strings
    (len > 1 simulates PowerPoint splitting a placeholder across runs)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    para = box.text_frame.paragraphs[0]
    for text in runs_text:
        run = para.add_run()
        run.text = text
    if notes is not None:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _render(prs, data, warnings):
    return Presentation(io.BytesIO(render_pptx(_bytes(prs), data, warnings)))


def _bytes(prs):
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _textbox_text(slide):
    return slide.shapes[0].text_frame.text


def test_simple_placeholder_resolves():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    warnings = []
    out = _render(prs, {"lead": {"contact_name": "Ala"}}, warnings)
    assert _textbox_text(out.slides[0]) == "Ala"
    assert warnings == []


def test_missing_value_becomes_empty_string_and_warns():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    warnings = []
    out = _render(prs, {"lead": {}}, warnings)
    assert _textbox_text(out.slides[0]) == ""
    assert any("lead.contact_name" in w for w in warnings)


def test_placeholder_split_across_runs_collapses_to_first_run():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.", "contact_name}}"])
    warnings = []
    out = _render(prs, {"lead": {"contact_name": "Ala"}}, warnings)
    para = out.slides[0].shapes[0].text_frame.paragraphs[0]
    assert [r.text for r in para.runs] == ["Ala", ""]


def test_repeat_participants_duplicates_slide_per_item():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    _add_slide(prs, ["{{participant.full_name}}"], notes="repeat:participants")
    warnings = []
    data = {
        "lead": {"contact_name": "Ala"},
        "participants": [{"full_name": "A"}, {"full_name": "B"}, {"full_name": "C"}],
    }
    out = _render(prs, data, warnings)
    slides = list(out.slides)
    assert len(slides) == 4  # 1 untouched + 3 repeats
    assert _textbox_text(slides[0]) == "Ala"
    assert [_textbox_text(s) for s in slides[1:]] == ["A", "B", "C"]
    assert warnings == []


def test_repeat_with_empty_list_drops_slide_and_warns():
    prs = _new_prs()
    _add_slide(prs, ["{{lead.contact_name}}"])
    _add_slide(prs, ["{{participant.full_name}}"], notes="repeat:participants")
    warnings = []
    data = {"lead": {"contact_name": "Ala"}, "participants": []}
    out = _render(prs, data, warnings)
    assert len(out.slides) == 1
    assert _textbox_text(out.slides[0]) == "Ala"
    assert any("repeat:participants" in w and "dropped" in w for w in warnings)


def test_repeat_testimonials_uses_testimonial_prefix():
    prs = _new_prs()
    _add_slide(prs, ["{{testimonial.client_name}}"], notes="repeat:testimonials")
    warnings = []
    data = {"testimonials": [{"client_name": "Firma X"}]}
    out = _render(prs, data, warnings)
    assert _textbox_text(list(out.slides)[0]) == "Firma X"
    assert warnings == []


def test_participant_nested_assessment_scores_resolve():
    prs = _new_prs()
    _add_slide(prs, ["{{participant.full_name}} {{participant.a.o}}"],
               notes="repeat:participants")
    warnings = []
    data = {"participants": [{"full_name": "Basia", "a": {"o": "B2"}}]}
    out = _render(prs, data, warnings)
    assert _textbox_text(list(out.slides)[0]) == "Basia B2"
    assert warnings == []


def test_participant_without_assessment_warns():
    prs = _new_prs()
    _add_slide(prs, ["{{participant.a.o}}"], notes="repeat:participants")
    warnings = []
    data = {"participants": [{"full_name": "Basia"}]}  # no "a" key
    _render(prs, data, warnings)
    assert any("participant.a.o" in w for w in warnings)
