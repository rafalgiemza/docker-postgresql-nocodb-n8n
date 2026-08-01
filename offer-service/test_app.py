"""Integration test for the /render endpoint - no NocoDB, no network.
app.py has zero NocoDB knowledge now, so this just drives the real FastAPI
handler with an in-memory template + data payload.

Run:  pytest offer-service/test_app.py -v
"""
import asyncio
import io
import json

from pptx import Presentation
from pptx.util import Inches
from starlette.datastructures import UploadFile

import app


def _template_pptx_bytes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.paragraphs[0].add_run().text = "{{lead.contact_name}}"

    repeat = prs.slides.add_slide(prs.slide_layouts[6])
    box2 = repeat.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box2.text_frame.paragraphs[0].add_run().text = "{{participant.full_name}}"
    repeat.notes_slide.notes_text_frame.text = "repeat:participant"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _call_render(template_bytes, data):
    upload = UploadFile(file=io.BytesIO(template_bytes), filename="template.pptx")
    response = asyncio.run(app.render(template=upload, data=json.dumps(data)))
    return response


def test_render_returns_rendered_pptx_and_no_warnings():
    data = {
        "lead": {"contact_name": "Ala"},
        "participant": [{"full_name": "Basia"}, {"full_name": "Czesiek"}],
    }
    response = _call_render(_template_pptx_bytes(), data)

    assert response.media_type == app.PPTX_MIME
    warnings = json.loads(response.headers["X-Warnings"])
    assert warnings == []

    rendered = Presentation(io.BytesIO(response.body))
    slides = list(rendered.slides)
    assert len(slides) == 3  # 1 untouched + 2 repeats
    assert slides[0].shapes[0].text_frame.text == "Ala"
    assert [s.shapes[0].text_frame.text for s in slides[1:]] == ["Basia", "Czesiek"]


def test_render_surfaces_warnings_in_header():
    data = {"lead": {}, "participant": []}
    response = _call_render(_template_pptx_bytes(), data)

    warnings = json.loads(response.headers["X-Warnings"])
    assert any("lead.contact_name" in w for w in warnings)
    assert any("repeat:participant" in w and "dropped" in w for w in warnings)
